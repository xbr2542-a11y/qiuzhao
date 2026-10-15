"""把之前导出的 Excel 总览重新导入岗位池（补回丢失的岗位）。

用途：导出的表格是一份岗位快照。如果之后岗位池被重抓/清理过，
可以用这个脚本把快照里"现在库里没有"的岗位补回去。

用法：
    python tools/import_export_xlsx.py "D:\秋招投递总览_xxx.xlsx" --dry-run
    python tools/import_export_xlsx.py "D:\秋招投递总览_xxx.xlsx" --ingest
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "app"))

import db  # noqa: E402
import ingest  # noqa: E402
import rules  # noqa: E402

# 导出表头 → 岗位字段
HEADER_MAP = {
    "公司": "company", "岗位": "title", "城市": "city", "岗位族": "family",
    "来源": "source", "匹配分": "match_score", "专业档位": "major_tier",
    "截止日期": "deadline", "剩余天数": None, "状态": None, "简历版本": None,
    "投递链接": "url", "备注": "note", "同链接": None,
}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("path")
    parser.add_argument("--ingest", action="store_true", help="真正写入岗位池")
    parser.add_argument("--dry-run", action="store_true", help="只看差异，不写库")
    args = parser.parse_args()

    from openpyxl import load_workbook

    db.init_db()
    workbook = load_workbook(args.path, read_only=True)
    sheet = workbook["投递总览"] if "投递总览" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(h).strip() if h else "" for h in rows[0]]
    print(f"读取 {args.path}")
    print(f"表头：{header}")
    print(f"数据行：{len(rows) - 1}")

    with db.connect() as conn:
        existing = {r["dedupe_key"] for r in conn.execute("SELECT dedupe_key FROM jobs").fetchall()}
        existing_urls = {r["url"] for r in conn.execute(
            "SELECT url FROM jobs WHERE url != ''").fetchall()}
    print(f"当前岗位池：{len(existing)} 条")

    missing: list[dict] = []
    for row in rows[1:]:
        record = {HEADER_MAP.get(header[i]): row[i] for i in range(min(len(header), len(row)))
                  if HEADER_MAP.get(header[i])}
        company = str(record.get("company") or "").strip()
        title = str(record.get("title") or "").strip()
        source = str(record.get("source") or "").strip()
        url = str(record.get("url") or "").strip()
        if not company or not title:
            continue
        key = rules.build_dedupe_key(company, title, source, url)
        if key in existing or (url and url in existing_urls):
            continue
        missing.append({
            "company": company, "title": title, "city": str(record.get("city") or ""),
            "source": source, "url": url, "deadline": str(record.get("deadline") or ""),
            "note": str(record.get("note") or ""),
            "family": record.get("family") or "",
        })

    print(f"\n快照里有、当前库里没有的岗位：{len(missing)} 条")
    for job in missing[:15]:
        print(f"  · {job['company'][:24]} | {job['title'][:28]} | {job['city'][:12]} | {job['source']}")
        print(f"      {job['url'][:96]}")
    if len(missing) > 15:
        print(f"  …… 其余 {len(missing) - 15} 条")

    if args.ingest and missing:
        result = ingest.upsert_jobs(missing, "历史导出补录")
        print(f"\n已补录：{result}")
    elif not args.ingest:
        print("\n（未写入，加 --ingest 才会补录）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
