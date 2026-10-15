"""导出 Excel。"""

from __future__ import annotations

from pathlib import Path

from db import EXPORT_DIR, FAMILIES, now
from rules import days_left


def export_workbook(jobs: list[dict], profile: dict) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def style_header(sheet, columns: list[tuple[str, int]]) -> None:
        for index, (title, width) in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index, value=title)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"

    summary = workbook.active
    summary.title = "投递总览"
    summary_columns = [("公司", 26), ("岗位", 30), ("城市", 10), ("岗位族", 10), ("来源", 14),
                       ("匹配分", 8), ("专业档位", 10), ("截止日期", 12), ("剩余天数", 10),
                       ("状态", 12), ("简历版本", 14), ("投递链接", 44), ("备注", 30)]
    style_header(summary, summary_columns)
    for row_index, job in enumerate(jobs, start=2):
        left = days_left(job.get("deadline", ""))
        values = [
            job.get("company", ""), job.get("title", ""), job.get("city", ""),
            FAMILIES.get(job.get("family", ""), job.get("family", "")), job.get("source", ""),
            job.get("match_score", 0), job.get("major_tier", 0), job.get("deadline", ""),
            "" if left is None else left, job.get("status", ""),
            job.get("resume_name", ""), job.get("url", ""), job.get("note", ""),
        ]
        for column_index, value in enumerate(values, start=1):
            summary.cell(row=row_index, column=column_index, value=value)

    details = workbook.create_sheet("筛选理由")
    detail_columns = [("公司", 26), ("岗位", 30), ("匹配分", 8), ("专业要求原文", 30), ("判定说明", 46)]
    style_header(details, detail_columns)
    for row_index, job in enumerate(jobs, start=2):
        details.cell(row=row_index, column=1, value=job.get("company", ""))
        details.cell(row=row_index, column=2, value=job.get("title", ""))
        details.cell(row=row_index, column=3, value=job.get("match_score", 0))
        details.cell(row=row_index, column=4, value=job.get("major_req", ""))
        details.cell(row=row_index, column=5, value=job.get("filter_note", ""))

    profile_sheet = workbook.create_sheet("我的档案")
    style_header(profile_sheet, [("字段", 20), ("内容", 60)])
    rows = [
        ("姓名", profile.get("full_name", "")), ("学校", profile.get("school", "")),
        ("专业", profile.get("major", "")), ("学历", profile.get("degree", "")),
        ("毕业时间", profile.get("grad_date", "")), ("政治面貌", profile.get("political", "")),
        ("意向城市", "、".join(profile.get("prefer_cities", []) or [])),
        ("手机", profile.get("phone", "")), ("邮箱", profile.get("email", "")),
        ("导出时间", now()),
    ]
    for row_index, (key, value) in enumerate(rows, start=2):
        profile_sheet.cell(row=row_index, column=1, value=key)
        profile_sheet.cell(row=row_index, column=2, value=value)

    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    stamp = now().replace(":", "").replace(" ", "_").replace("-", "")
    target = EXPORT_DIR / f"秋招投递总览_{stamp}.xlsx"
    workbook.save(target)
    return target
