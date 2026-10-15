"""简历文件保存与结构化解析。"""

from __future__ import annotations

import base64
import hashlib
import re
from pathlib import Path
from typing import Any

from db import RESUME_DIR, now


def _safe_name(name: str) -> str:
    cleaned = re.sub(r"[\\/:*?\"<>|]+", "_", name).strip() or "resume"
    return cleaned[:120]


def save_file(filename: str, data_base64: str) -> dict[str, Any]:
    raw = base64.b64decode(data_base64)
    digest = hashlib.sha256(raw).hexdigest()[:16]
    suffix = Path(filename).suffix.lower() or ".bin"
    target = RESUME_DIR / f"{digest}{suffix}"
    target.write_bytes(raw)
    return {
        "filename": filename,
        "stored_path": str(target),
        "size": len(raw),
        "sha256": digest,
    }


def read_text(path: str) -> str:
    p = Path(path)
    if not p.exists():
        return ""
    suffix = p.suffix.lower()
    try:
        if suffix == ".pdf":
            import pypdf

            reader = pypdf.PdfReader(str(p))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        if suffix in {".docx", ".doc"}:
            import docx

            document = docx.Document(str(p))
            parts = [para.text for para in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            return "\n".join(parts)
        if suffix in {".txt", ".md"}:
            return p.read_text(encoding="utf-8", errors="ignore")
        if suffix == ".html":
            return re.sub(r"<[^>]+>", " ", p.read_text(encoding="utf-8", errors="ignore"))
    except Exception as exc:  # 解析失败不应该阻塞上传
        return f"__PARSE_ERROR__ {exc}"
    return ""


LABELS = {
    "full_name": ["姓名", "名字"],
    "school": ["学校", "毕业院校", "院校", "就读院校"],
    "major": ["专业", "所学专业", "专业名称"],
    "degree": ["学历", "学位", "最高学历"],
    "grad_date": ["毕业时间", "毕业年月", "预计毕业", "毕业日期"],
    "phone": ["电话", "手机", "联系电话", "手机号", "联系方式"],
    "email": ["邮箱", "电子邮件", "Email", "E-mail"],
    "political": ["政治面貌", "政治", "党员"],
    "hometown": ["籍贯", "户籍", "生源地", "家庭住址"],
}


def _grab_labeled(text: str, labels: list[str]) -> str:
    for label in labels:
        pattern = rf"{re.escape(label)}\s*[:：]?\s*([^\n\r|｜]{{1,40}})"
        match = re.search(pattern, text)
        if match:
            value = match.group(1).strip(" |｜·,，。;；")
            if value:
                return value
    return ""


def parse_resume_text(text: str) -> dict[str, Any]:
    result: dict[str, Any] = {"raw_length": len(text or ""), "hints": []}
    if not text or text.startswith("__PARSE_ERROR__"):
        result["hints"].append("未能提取到文本，可能是扫描件或图片版简历，需要人工补填字段")
        return result

    for field, labels in LABELS.items():
        value = _grab_labeled(text, labels)
        if value:
            result[field] = value

    phone = re.search(r"1[3-9]\d{9}", text)
    if phone:
        result["phone"] = phone.group(0)
    email = re.search(r"[\w.\-+]+@[\w\-]+\.[A-Za-z]{2,}", text)
    if email:
        result["email"] = email.group(0)

    if not result.get("degree"):
        for name in ["博士", "硕士", "研究生", "本科", "学士", "大专", "专科"]:
            if name in text:
                result["degree"] = name
                break
    if not result.get("major"):
        match = re.search(r"([\u4e00-\u9fa5]{2,12}(?:设计|工程|管理|科学|技术|经济|贸易|传播|教育|学))\s*(?:专业)?", text)
        if match:
            result["major"] = match.group(1)
            result["hints"].append("专业为猜测值，请在档案里确认")

    if not result.get("school"):
        match = re.search(r"([\u4e00-\u9fa5]{2,15}(?:大学|学院|职业技术学院))", text)
        if match:
            result["school"] = match.group(1)

    result["skills"] = [k for k in ["Excel", "SQL", "Python", "Tableau", "Power BI", "SPSS", "英语",
                                     "CET-4", "CET-6", "雅思", "托福", "Photoshop", "CAD", "SketchUp"]
                            if k.lower() in text.lower()]
    if not result.get("hints"):
        result["hints"] = ["字段为自动提取结果，投递前请核对"]
    return result


def parse_resume_file(path: str) -> dict[str, Any]:
    parsed = parse_resume_text(read_text(path))
    parsed["parsed_at"] = now()
    return parsed


# ---------------------------------------------------------------- 结构化解析

SECTION_ALIASES = {
    "教育背景": "education", "教育经历": "education", "学历背景": "education",
    "服役经历": "service", "部队经历": "service", "军旅经历": "service",
    "校园经历": "campus", "社团经历": "campus", "实践经历": "campus", "在校经历": "campus",
    "实习经历": "intern", "工作经历": "work", "项目经历": "project", "项目经验": "project",
    "荣誉奖项": "honor", "获奖情况": "honor", "荣誉": "honor",
    "技能证书": "skill", "技能": "skill", "证书": "skill", "专业技能": "skill",
    "自我评价": "summary", "个人评价": "summary", "自我介绍": "summary",
}

SECTION_LABELS = {
    "education": "教育", "service": "服役", "campus": "校园", "intern": "实习",
    "work": "工作", "project": "项目", "honor": "荣誉", "skill": "技能",
    "summary": "自述", "header": "基本信息",
}

DEGREE_RANK = {"专科": 2, "大专": 2, "本科": 3, "学士": 3, "硕士": 4, "研究生": 4, "博士": 5}


def read_docx_sections(path: str) -> dict[str, list[str]]:
    """按文档真实顺序把段落/表格归到章节下（简历常用表格排版，必须逐格读）。"""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(path)
    sections: dict[str, list[str]] = {}
    current = "header"

    def push(text: str) -> None:
        text = (text or "").strip()
        if text:
            sections.setdefault(current, []).append(text)

    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            text = Paragraph(child, document).text.strip()
            if not text:
                continue
            key = SECTION_ALIASES.get(text.strip("：: "))
            if key:
                current = key
                continue
            push(text)
        elif tag == "tbl":
            table = Table(child, document)
            cells: list[str] = []
            for row in table.rows:
                for cell in row.cells:
                    value = cell.text.strip()
                    if value and (not cells or cells[-1] != value):
                        cells.append(value)
            if not cells:
                continue
            key = SECTION_ALIASES.get(cells[0].strip("：: "))
            if key:
                current = key
                cells = cells[1:]
            for value in cells:
                push(value)
    return sections


def extract_profile_fields(sections: dict[str, list[str]]) -> dict[str, Any]:
    """从章节文本里抽取档案字段：只填有把握的，拿不准的留空由用户确认。"""
    flat = "\n".join(line for lines in sections.values() for line in lines)
    profile: dict[str, Any] = {}

    # 姓名：基本信息里第一行通常是姓名（2-4 个汉字，且不是标签）
    for line in sections.get("header", []):
        candidate = line.strip()
        if 2 <= len(candidate) <= 4 and re.fullmatch(r"[\u4e00-\u9fa5·]{2,4}", candidate) \
                and not any(k in candidate for k in ["意向", "求职", "电话", "邮箱", "专业", "大学", "学院"]):
            profile["full_name"] = candidate
            break

    phone = re.search(r"1[3-9]\d{9}", flat)
    if phone:
        profile["phone"] = phone.group(0)
    email = re.search(r"[\w.\-+]+@[\w\-]+\.[A-Za-z]{2,}", flat)
    if email:
        profile["email"] = email.group(0)

    # 求职意向
    intent = re.search(r"求职意向[：:]\s*([^\n]{2,80})", flat)
    if intent:
        profile["job_intent"] = intent.group(1).strip()

    # 现居城市
    city = re.search(r"现居(?:于)?([\u4e00-\u9fa5]{2,8}?[市省])", flat)
    if city:
        profile["current_city"] = city.group(1)

    # 教育：取学历最高的一条作为当前学历
    best: tuple[int, str, str, str] = (0, "", "", "")
    for line in sections.get("education", []):
        match_degree = next((d for d in DEGREE_RANK if d in line), "")
        match_school = re.search(r"([\u4e00-\u9fa5]{2,15}(?:大学|学院|职业学院|职业技术学院))", line)
        match_major = re.search(r"([\u4e00-\u9fa5]{2,12}(?:设计|工程|管理|科学|技术|经济|贸易|传播|教育|学))(?![校对])", line)
        match_date = re.search(r"(\d{4}\.\d{1,2}\s*[-—~至]\s*(?:\d{4}\.\d{1,2}|至今))", line)
        if not match_school:
            continue
        rank = DEGREE_RANK.get(match_degree, 0)
        if rank >= best[0]:
            best = (rank, match_school.group(1),
                    match_major.group(1) if match_major else "",
                    match_date.group(1) if match_date else "")
    if best[1]:
        profile["school"] = best[1]
        profile["major"] = best[2]
        profile["degree"] = next((d for d in DEGREE_RANK if DEGREE_RANK[d] == best[0]), "")
        if best[3]:
            profile["education_period"] = best[3].replace(" ", "")

    # 退役大学生士兵
    if any(k in flat for k in ["义务兵", "服兵役", "部队", "武警", "解放军", "参军", "退役军人"]):
        profile["veteran"] = 1

    # 技能
    skill_blob = "\n".join(sections.get("skill", []))
    known_skills = ["Word", "Excel", "PowerPoint", "CAD", "SketchUp", "SU", "V-Ray", "VRay",
                    "Photoshop", "PS", "AI", "SQL", "Python", "Tableau", "Power BI", "SPSS",
                    "Office", "WPS", "AI工具", "ChatGPT", "DeepSeek", "Codex", "英语", "CET-4",
                    "CET-6", "雅思", "托福", "普通话"]
    profile["skills"] = [s for s in known_skills if s.lower() in skill_blob.lower()]
    return profile


def build_facts(resume_id: int, sections: dict[str, list[str]], parsed: dict[str, Any]) -> list[dict[str, Any]]:
    """把简历内容拆成可引用的事实条目——AI 生成时只能引用这里的东西，杜绝编造。"""
    facts: list[dict[str, Any]] = []
    counter = 1
    for key, lines in sections.items():
        if key == "header":
            continue
        category = SECTION_LABELS.get(key, key)
        for line in lines:
            text = line.strip("•·-— \t")
            if not text or len(text) < 2:
                continue
            org = ""
            school = re.search(r"([\u4e00-\u9fa5]{2,15}(?:大学|学院|部队|公司|集团|协会|研究会|社))", text)
            if school:
                org = school.group(1)
            period = re.search(r"(\d{4}\.\d{1,2}\s*[-—~至]\s*(?:\d{4}\.\d{1,2}|至今))", text)
            facts.append({
                "resume_id": resume_id,
                "fact_id": f"F{counter:02d}",
                "category": category,
                "title": text[:60],
                "org": org,
                "role": "",
                "start_date": (period.group(1).split("-")[0].strip() if period else ""),
                "end_date": (period.group(1).split("-")[-1].strip() if period else ""),
                "detail": text,
                "keywords": parsed.get("skills", [])[:6] if key == "skill" else [],
                "confidence": 0.7,
                "confirmed": 0,
            })
            counter += 1
    return facts
